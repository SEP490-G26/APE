import csv
import json
import os
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
APP_DATA = REPO / "src-dotnet" / "Ape.AiModule.Api" / "App_Data"
HISTORY_ROOT = APP_DATA / "ai-config-history"
RUN_HISTORY_ROOT = APP_DATA / "run-history"
REPORT_DIR = REPO / "Docs" / "Prototype_Document" / "Guides" / "Templates" / "AI_Report"
TEST_REPORT_DIR = REPO / "Test"
WORKBOOK_TEMPLATE_PATH = REPORT_DIR / "AI_Benchmark_Report_Template.xlsx"
WORKBOOK_PATH = TEST_REPORT_DIR / "AI_Benchmark_Report.xlsx"
CHANGELOG_CSV = TEST_REPORT_DIR / "ai_config_version_changelog.csv"
RUN_CONFIG_CSV = TEST_REPORT_DIR / "run_config_references.csv"
RUN_INDEX_CSV = TEST_REPORT_DIR / "run_history_index_export.csv"
INPUT_EVIDENCE_CSV = TEST_REPORT_DIR / "input_quality_evidence.csv"
QUESTION_EVIDENCE_CSV = TEST_REPORT_DIR / "question_generation_evidence.csv"
MENTOR_EVIDENCE_CSV = TEST_REPORT_DIR / "code_mentor_evidence.csv"
INPUT_BENCHMARK_CSV = TEST_REPORT_DIR / "input_quality_benchmark.csv"
QUESTION_BENCHMARK_CSV = TEST_REPORT_DIR / "question_generation_benchmark.csv"
MENTOR_BENCHMARK_CSV = TEST_REPORT_DIR / "code_mentor_benchmark.csv"
EVIDENCE_ROOT = TEST_REPORT_DIR / "report_evidence"
DEFAULT_TESTER = os.environ.get("AI_BENCHMARK_TESTER") or os.environ.get("USERNAME") or os.environ.get("USER") or ""

CURRENT_FILES = {
    "prompts": APP_DATA / "ai-prompts.json",
    "policies": [
        APP_DATA / "gatekeeper-policy.json",
        APP_DATA / "extracted-content-policy.json",
        APP_DATA / "embedding-tagging-policy.json",
        APP_DATA / "code-mentor-policy.json",
    ],
    "rubrics": APP_DATA / "ai-rubrics",
    "groundtruth": APP_DATA / "ai-groundtruth",
}

TAB_COLORS = {
    "AI Config Version Tracking": "1F4E78",
    "Prompt Version Comparison": "8E5C2A",
    "Run History Index": "355E3B",
    "Run Config References": "5A4B81",
    "Input Quality Benchmark": "4E7C67",
    "Question Generation Benchmark": "8B5E34",
    "Code Mentor Benchmark": "5D4383",
    "Input Quality Evidence": "7A9E7E",
    "Question Generation Evidence": "A56A43",
    "Code Mentor Evidence": "6A4C93",
}


def bootstrap_output_structure():
    TEST_REPORT_DIR.mkdir(parents=True, exist_ok=True)
    EVIDENCE_ROOT.mkdir(parents=True, exist_ok=True)
    for child in ("input-quality", "question-generation", "code-mentor"):
        (EVIDENCE_ROOT / child).mkdir(parents=True, exist_ok=True)


def ensure_required_paths():
    required_dirs = {
        "repo": REPO,
        "app data": APP_DATA,
        "report template dir": REPORT_DIR,
        "run history": RUN_HISTORY_ROOT,
    }
    missing_dirs = [f"{label}: {path}" for label, path in required_dirs.items() if not path.exists()]
    if missing_dirs:
        raise FileNotFoundError("Missing required directories:\n" + "\n".join(missing_dirs))

    required_files = [WORKBOOK_TEMPLATE_PATH, CURRENT_FILES["prompts"]]
    missing_files = [str(path) for path in required_files if not path.exists()]
    if missing_files:
        raise FileNotFoundError("Missing required files:\n" + "\n".join(missing_files))


def load_json(path: Path):
    with path.open("r", encoding="utf-8-sig") as handle:
        return json.load(handle)


def normalize_text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def collapse_ws(value):
    return " ".join(normalize_text(value).split())


def shorten(value, limit=280):
    text = collapse_ws(value)
    if len(text) <= limit:
        return text
    return text[: limit - 3].rstrip() + "..."


def join_values(values, sep=" | ", limit=8):
    if not values:
        return ""
    flat = [collapse_ws(value) for value in values if collapse_ws(value)]
    if not flat:
        return ""
    if len(flat) > limit:
        return sep.join(flat[:limit]) + f" | +{len(flat) - limit} more"
    return sep.join(flat)


def parse_dt(value):
    if not value:
        return datetime.min
    text = str(value).replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return datetime.min


def normalize_number(value, digits=None):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return round(float(value), digits) if digits is not None else float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    return round(number, digits) if digits is not None else number


def normalize_int(value):
    number = normalize_number(value)
    if number is None:
        return None
    return int(round(number))


def yes_no(value):
    if value in (None, ""):
        return ""
    return "yes" if bool(value) else "no"


def safe_ratio(numerator, denominator, digits=4):
    left = normalize_number(numerator)
    right = normalize_number(denominator)
    if left is None or right in (None, 0):
        return None
    return round(left / right, digits)


def score_cost_index(score, cost_usd):
    left = normalize_number(score)
    right = normalize_number(cost_usd)
    if left is None or right is None or right <= 0:
        return None
    return round(left / right, 6)


def score_latency_index(score, latency_ms):
    left = normalize_number(score)
    right = normalize_number(latency_ms)
    if left is None or right is None or right <= 0:
        return None
    return round(left / right, 6)


def estimate_token_count(text):
    content = normalize_text(text)
    if not content:
        return 0
    # Pragmatic estimate for report autofill only.
    return max(1, round(len(content) / 4))


def get_config_ref_map(config_refs):
    return {ref["artifact_type"]: ref for ref in config_refs or []}


def get_config_version(config_refs, artifact_type, default="N/A"):
    for ref in config_refs or []:
        if ref.get("artifact_type") == artifact_type:
            stable_id = normalize_text(ref.get("stable_id"))
            version = normalize_text(ref.get("version"))
            if stable_id or version:
                return f"{stable_id}:{version}" if stable_id else version
    return default


def get_stage_cost(response, stage_name_prefix):
    for stage in response.get("StageLogs", []) or []:
        stage_name = normalize_text(stage.get("StageName"))
        if stage_name.startswith(stage_name_prefix):
            return stage.get("Cost", {}) or {}
    return {}


def get_stage_costs(response, stage_name_prefix):
    return [stage.get("Cost", {}) or {} for stage in response.get("StageLogs", []) or [] if normalize_text(stage.get("StageName")).startswith(stage_name_prefix)]


def ensure_sheet(workbook, title):
    if title in workbook.sheetnames:
        return workbook[title]
    return workbook.create_sheet(title)


def clear_data_rows(worksheet):
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)


def write_rows(worksheet, headers, rows):
    for col_idx, header in enumerate(headers, start=1):
        worksheet.cell(1, col_idx, header)
    clear_data_rows(worksheet)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, header in enumerate(headers, start=1):
            worksheet.cell(row_idx, col_idx, row.get(header, ""))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, worksheet.max_row)}"
    for col_idx, header in enumerate(headers, start=1):
        samples = [len(str(header))]
        samples.extend(len(str(row.get(header, ""))) for row in rows[:50])
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(samples) + 2, 48)


def set_tab_style(worksheet):
    if worksheet.title in TAB_COLORS:
        worksheet.sheet_properties.tabColor = TAB_COLORS[worksheet.title]
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)


def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def write_text(path, content):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def format_key_value_lines(pairs):
    return "\n".join(f"- {label}: {value}" for label, value in pairs if collapse_ws(value))


def format_question_block(question, index):
    qtype = normalize_text(question.get("Type") or question.get("type"))
    lines = [
        f"Question {index}: {collapse_ws(question.get('Title') or question.get('title'))}",
        f"Type: {qtype}",
        f"Difficulty: {collapse_ws(question.get('Difficulty') or question.get('difficulty'))}",
    ]
    description = normalize_text(question.get("Description") or question.get("description"))
    if description:
        lines.append(f"Description: {collapse_ws(description)}")
    tags = question.get("TopicTags") or question.get("topic_tags") or []
    if tags:
        lines.append(f"Topic tags: {join_values(tags, sep=', ')}")
    sources = question.get("SourceChunkIds") or question.get("source_chunk_ids") or []
    if sources:
        lines.append(f"Source chunks: {join_values(sources, sep=', ')}")
    options = question.get("Options") or question.get("options") or []
    if options:
        lines.append("Options:")
        lines.extend(f"  {collapse_ws(option)}" for option in options)
    answers = question.get("CorrectAnswer") or question.get("correct_answer") or []
    if answers:
        lines.append(f"Correct answer: {join_values(answers, sep=', ')}")
    explanation = normalize_text(question.get("Explanation") or question.get("explanation"))
    if explanation:
        lines.append(f"Explanation: {collapse_ws(explanation)}")
    skeleton = question.get("SkeletonCode") or question.get("skeleton_code") or []
    if skeleton:
        lines.append("Skeleton code:")
        for item in skeleton:
            lines.append(
                f"  {collapse_ws(item.get('Filename') or item.get('filename'))}: "
                f"{shorten(item.get('Content') or item.get('content'), 180)}"
            )
    solution = question.get("SolutionCode") or question.get("solution_code") or []
    if solution:
        lines.append("Solution code:")
        for item in solution:
            lines.append(
                f"  {collapse_ws(item.get('Filename') or item.get('filename'))}: "
                f"{shorten(item.get('Content') or item.get('content'), 180)}"
            )
    tests = question.get("TestCases") or question.get("test_cases") or []
    if tests:
        lines.append("Test cases:")
        for test in tests[:5]:
            lines.append(
                "  input="
                + collapse_ws(test.get("Input") or test.get("input"))
                + "; expected="
                + collapse_ws(test.get("ExpectedOutput") or test.get("expected_output"))
                + "; hidden="
                + normalize_text(test.get("IsHidden") if "IsHidden" in test else test.get("is_hidden"))
            )
        if len(tests) > 5:
            lines.append(f"  +{len(tests) - 5} more test cases")
    return "\n".join(lines)


def load_snapshot_index():
    rows = []
    by_key = {}
    if not HISTORY_ROOT.exists():
        return rows, by_key
    for path in sorted(HISTORY_ROOT.rglob("*.json")):
        payload = load_json(path)
        row = {
            "artifact_type": normalize_text(payload.get("artifact_type")),
            "stable_id": normalize_text(payload.get("stable_id")),
            "version": normalize_text(payload.get("version")),
            "function_name": normalize_text(payload.get("function_name")),
            "subject_code": normalize_text(payload.get("subject_code")),
            "question_type": normalize_text(payload.get("question_type")),
            "language": normalize_text(payload.get("language")),
            "captured_at": normalize_text(payload.get("captured_at")),
            "source": normalize_text(payload.get("source")),
            "source_commit": normalize_text(payload.get("source_commit")),
            "captured_from": normalize_text(payload.get("captured_from")),
            "change_reason": normalize_text(payload.get("change_reason")),
            "change_summary": normalize_text(payload.get("change_summary")),
            "is_active": normalize_text(payload.get("is_active")),
            "content_hash": normalize_text(payload.get("content_hash")),
            "snapshot_file": str(path.relative_to(REPO)).replace("\\", "/"),
        }
        rows.append(row)
        by_key[(row["artifact_type"], row["stable_id"], row["version"], row["content_hash"])] = row
    rows.sort(key=lambda item: (item["artifact_type"], item["stable_id"], parse_dt(item["captured_at"])))
    return rows, by_key


def load_current_catalog():
    rows = []
    prompts = load_json(CURRENT_FILES["prompts"])
    for item in prompts:
        rows.append(
            {
                "config_type": "prompt",
                "config_key": normalize_text(item.get("Key")),
                "version": normalize_text(item.get("Version")),
                "name_or_description": normalize_text(item.get("Description")),
                "subject_code": "",
                "question_type": "",
                "status": "active" if item.get("IsActive") else "inactive",
                "system_prompt_len": len(normalize_text(item.get("SystemPrompt"))),
                "user_prompt_len": len(normalize_text(item.get("UserPrompt"))),
                "updated_at": normalize_text(item.get("UpdatedAt")),
                "source_file": "src-dotnet/Ape.AiModule.Api/App_Data/ai-prompts.json",
            }
        )
    for path in CURRENT_FILES["policies"]:
        item = load_json(path)
        rows.append(
            {
                "config_type": "policy",
                "config_key": normalize_text(item.get("policy_id") or path.stem),
                "version": normalize_text(item.get("version") or "v1"),
                "name_or_description": normalize_text(item.get("description") or item.get("function_name") or path.stem),
                "subject_code": normalize_text(item.get("subject_code")),
                "question_type": normalize_text(item.get("question_type")),
                "status": "active" if item.get("is_active", True) else "inactive",
                "system_prompt_len": "",
                "user_prompt_len": "",
                "updated_at": normalize_text(item.get("updated_at")),
                "source_file": str(path.relative_to(REPO)).replace("\\", "/"),
            }
        )
    for folder_name, config_type, id_field in [("ai-rubrics", "rubric", "rubric_id"), ("ai-groundtruth", "groundtruth", "dataset_id")]:
        folder = APP_DATA / folder_name
        for path in sorted(folder.glob("*.json")):
            if path.name == "README.json":
                continue
            item = load_json(path)
            rows.append(
                {
                    "config_type": config_type,
                    "config_key": normalize_text(item.get(id_field) or path.stem),
                    "version": normalize_text(item.get("version") or "v1"),
                    "name_or_description": normalize_text(item.get("title") or item.get("description") or item.get("function_name") or path.stem),
                    "subject_code": normalize_text(item.get("subject_code")),
                    "question_type": normalize_text(item.get("question_type")),
                    "status": "active" if item.get("is_active", True) else "inactive",
                    "system_prompt_len": "",
                    "user_prompt_len": "",
                    "updated_at": normalize_text(item.get("updated_at")),
                    "source_file": str(path.relative_to(REPO)).replace("\\", "/"),
                }
            )
    rows.sort(key=lambda item: (item["config_type"], item["config_key"]))
    return rows


def load_active_prompt_versions():
    active = {}
    prompts = load_json(CURRENT_FILES["prompts"])
    for item in prompts:
        if item.get("IsActive"):
            key = normalize_text(item.get("Key"))
            version = normalize_text(item.get("Version"))
            if key:
                active[key] = version
    return active


def load_changelog_rows():
    with CHANGELOG_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    rows.sort(key=lambda item: (item["artifact_type"], item["stable_id"], parse_dt(item["commit_date_utc"])))
    return rows


def build_version_tracking_rows(changelog_rows):
    grouped = defaultdict(list)
    for row in changelog_rows:
        grouped[(row["artifact_type"], row["stable_id"])].append(row)
    results = []
    for (_, stable_id), items in sorted(grouped.items()):
        parent_version_id = ""
        for idx, item in enumerate(items):
            version_id = f"{item['stable_id']}::{item['version']}::{item['content_hash'][:12]}"
            results.append(
                {
                    "config_type": item["artifact_type"][:-1] if item["artifact_type"].endswith("s") else item["artifact_type"],
                    "function_name": item.get("function_name", ""),
                    "component_name": stable_id,
                    "version_id": version_id,
                    "parent_version_id": parent_version_id,
                    "status": "active" if idx == len(items) - 1 else "historical",
                    "applied_to_models": "",
                    "subject_code": item.get("subject_code", ""),
                    "question_type": item.get("question_type", ""),
                    "effective_from": item.get("commit_date_utc", ""),
                    "effective_to": "",
                    "change_reason": item.get("change_reason", ""),
                    "change_summary": item.get("change_summary", ""),
                    "expected_impact": "",
                    "dataset_used": "",
                    "benchmark_run_range": "",
                    "quality_before": "",
                    "quality_after": "",
                    "cost_before": "",
                    "cost_after": "",
                    "stability_before": "",
                    "stability_after": "",
                    "owner_note": "",
                    "source_commit": item.get("source_commit", ""),
                    "snapshot_path": item.get("snapshot_path", ""),
                }
            )
            parent_version_id = version_id
    return results


def build_prompt_comparison_rows(snapshot_rows):
    grouped = defaultdict(list)
    for row in snapshot_rows:
        if row["artifact_type"] == "prompts":
            grouped[row["stable_id"]].append(row)
    current_prompts = {item["Key"]: item for item in load_json(CURRENT_FILES["prompts"])}
    results = []
    for prompt_key, items in sorted(grouped.items()):
        items.sort(key=lambda item: parse_dt(item["captured_at"]))
        current = items[-1]
        previous = items[-2] if len(items) > 1 else None
        current_item = current_prompts.get(prompt_key, {})
        previous_payload = load_json(REPO / previous["snapshot_file"]) if previous else {}
        previous_content = previous_payload.get("content", {}) if previous else {}
        results.append(
            {
                "prompt_key": prompt_key,
                "current_version": current.get("version", ""),
                "previous_version": previous.get("version", "") if previous else "",
                "description": current.get("change_reason", "") or current.get("stable_id", ""),
                "current_system_len": len(normalize_text(current_item.get("SystemPrompt"))),
                "previous_system_len": len(normalize_text(previous_content.get("system_prompt"))),
                "current_user_len": len(normalize_text(current_item.get("UserPrompt"))),
                "previous_user_len": len(normalize_text(previous_content.get("user_prompt"))),
                "change_detected": "yes" if previous else "no",
                "change_reason": current.get("change_reason", ""),
                "dataset_used": "",
                "benchmark_run_ids_before": "",
                "benchmark_run_ids_after": "",
                "quality_before": "",
                "quality_after": "",
                "cost_before": "",
                "cost_after": "",
                "decision_note": current.get("change_summary", ""),
                "source_commit_current": current.get("source_commit", ""),
                "source_commit_previous": previous.get("source_commit", "") if previous else "",
                "current_hash": current.get("content_hash", ""),
                "previous_hash": previous.get("content_hash", "") if previous else "",
            }
        )
    return results


def collect_runs(snapshot_index):
    run_rows = []
    config_ref_rows = []
    config_refs_by_run = defaultdict(list)
    run_payloads = []
    for path in sorted(RUN_HISTORY_ROOT.rglob("*.json")):
        payload = load_json(path)
        summary = payload.get("Summary", {})
        request = payload.get("Request", {})
        response = payload.get("Response", {})
        run_id = normalize_text(summary.get("RunId"))
        run_row = {
            "run_id": run_id,
            "function_name": normalize_text(summary.get("FunctionName")),
            "route_key": normalize_text(summary.get("RouteKey")),
            "recorded_at": normalize_text(summary.get("RecordedAt")),
            "status": normalize_text(summary.get("Status")),
            "model_summary": normalize_text(summary.get("ModelSummary")),
            "subject": normalize_text(request.get("Subject")),
            "question_type": normalize_text(request.get("QuestionType")),
            "difficulty": normalize_text(request.get("Difficulty")),
            "input_tokens": normalize_text(summary.get("Totals", {}).get("InputTokens")),
            "output_tokens": normalize_text(summary.get("Totals", {}).get("OutputTokens")),
            "total_cost_usd": normalize_text(summary.get("Totals", {}).get("TotalCostUsd")),
            "latency_ms": normalize_text(summary.get("Totals", {}).get("LatencyMs")),
            "prompt_key_guess": "",
            "prompt_version_ref": "",
            "config_link_mode": "config_refs" if summary.get("ConfigReferences") else "none",
            "history_file": str(path.relative_to(REPO)).replace("\\", "/"),
            "usage_source": normalize_text(summary.get("UsageSource")),
            "error": normalize_text(summary.get("Error")),
        }
        run_rows.append(run_row)
        run_payloads.append(
            {
                "path": path,
                "payload": payload,
                "summary": summary,
                "request": request,
                "response": response,
                "run_row": run_row,
            }
        )
        for ref in summary.get("ConfigReferences") or []:
            artifact_type = normalize_text(ref.get("ArtifactType"))
            stable_id = normalize_text(ref.get("StableId"))
            version = normalize_text(ref.get("Version"))
            content_hash = normalize_text(ref.get("ContentHash"))
            candidates = [
                (artifact_type + "s" if artifact_type in {"prompt", "policy", "rubric"} else artifact_type, stable_id, version, content_hash),
                (artifact_type if artifact_type.endswith("s") else artifact_type, stable_id, version, content_hash),
            ]
            snapshot = None
            for key in candidates:
                snapshot = snapshot_index.get(key)
                if snapshot:
                    break
            config_ref_row = {
                "run_id": run_id,
                "function_name": normalize_text(summary.get("FunctionName")),
                "route_key": normalize_text(summary.get("RouteKey")),
                "recorded_at": normalize_text(summary.get("RecordedAt")),
                "artifact_type": artifact_type,
                "stable_id": stable_id,
                "version": version,
                "subject_code": normalize_text(ref.get("SubjectCode")),
                "question_type": normalize_text(ref.get("QuestionType")),
                "language": normalize_text(ref.get("Language")),
                "content_hash": content_hash,
                "source": normalize_text(ref.get("Source")),
                "source_commit": snapshot.get("source_commit", "") if snapshot else "",
                "change_reason": snapshot.get("change_reason", "") if snapshot else "",
                "change_summary": snapshot.get("change_summary", "") if snapshot else "",
                "snapshot_file": snapshot.get("snapshot_file", "") if snapshot else "",
                "history_file": str(path.relative_to(REPO)).replace("\\", "/"),
            }
            config_ref_rows.append(config_ref_row)
            config_refs_by_run[run_id].append(config_ref_row)
    run_rows.sort(key=lambda item: parse_dt(item["recorded_at"]))
    config_ref_rows.sort(key=lambda item: (parse_dt(item["recorded_at"]), item["run_id"], item["artifact_type"], item["stable_id"]))
    for row in run_rows:
        refs = [ref for ref in config_refs_by_run.get(row["run_id"], []) if ref["artifact_type"] == "prompt"]
        row["prompt_version_ref"] = " | ".join(f"{ref['stable_id']}:{ref['version']}" for ref in refs)
        row["prompt_key_guess"] = " | ".join(ref["stable_id"] for ref in refs)
    return run_rows, config_ref_rows, config_refs_by_run, run_payloads


def build_question_markdown(run, evidence_rel_path):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    questions = response.get("Questions", []) or []
    review = response.get("Review", {}) or {}
    schema = response.get("SchemaMapping", {}) or {}
    difficulty = response.get("DifficultyAlignment", {}) or {}
    chunks = request.get("Chunks", []) or []
    sections = [chunk.get("SectionTitle") for chunk in chunks if chunk.get("SectionTitle")]
    md = []
    md.append(f"# Question Generation Evidence: {summary.get('RunId', '')}")
    md.append("")
    md.append("## Run Metadata")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Recorded at", summary.get("RecordedAt")),
                ("Status", summary.get("Status")),
                ("Subject", request.get("Subject")),
                ("Question type", request.get("QuestionType")),
                ("Difficulty", request.get("Difficulty")),
                ("Mode", request.get("Mode")),
                ("Generator model", request.get("GeneratorModel")),
                ("Reviewer model", request.get("ReviewerModel")),
                ("Attempts used", response.get("AttemptsUsed")),
                ("Total cost USD", summary.get("Totals", {}).get("TotalCostUsd")),
                ("Latency ms", summary.get("Totals", {}).get("LatencyMs")),
                ("Evidence file", evidence_rel_path),
            ]
        )
    )
    md.append("")
    md.append("## Input Context")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Chunk count", len(chunks)),
                ("Section titles", join_values(sections, sep=", ")),
                ("Chunk ids", join_values([chunk.get("ChunkId") for chunk in chunks], sep=", ", limit=12)),
                ("Topic tags", join_values([tag for chunk in chunks for tag in (chunk.get("TopicTags") or [])], sep=", ", limit=15)),
            ]
        )
    )
    md.append("")
    if chunks:
        md.append("### Context Preview")
        md.append("")
        for idx, chunk in enumerate(chunks[:3], start=1):
            md.append(f"Chunk {idx}: {collapse_ws(chunk.get('SectionTitle') or chunk.get('ChunkId'))}")
            md.append(shorten(chunk.get("ContentText"), 1000))
            md.append("")
    md.append("## Generated Questions")
    md.append("")
    if questions:
        for idx, question in enumerate(questions, start=1):
            md.append(format_question_block(question, idx))
            md.append("")
    else:
        md.append("No generated questions found.")
        md.append("")
    md.append("## Review Result")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Review status", review.get("ReviewStatus")),
                ("Score", review.get("Score")),
                ("Schema valid", review.get("SchemaValid")),
                ("Content grounded", review.get("ContentGrounded")),
                ("Needs revision", review.get("NeedsRevision")),
                ("Reviewer model", review.get("ReviewerModel")),
                ("Estimated difficulty", difficulty.get("EstimatedDifficulty")),
                ("Difficulty aligned", difficulty.get("IsAligned")),
            ]
        )
    )
    md.append("")
    if review.get("Issues"):
        md.append("Issues:")
        md.extend(f"- {collapse_ws(item)}" for item in review.get("Issues", []))
        md.append("")
    if review.get("Suggestions"):
        md.append("Suggestions:")
        md.extend(f"- {collapse_ws(item)}" for item in review.get("Suggestions", []))
        md.append("")
    if schema.get("Issues"):
        md.append("Rule-based / schema issues:")
        for issue in schema.get("Issues", []):
            label = f"{issue.get('Code')}: {issue.get('Message')}"
            md.append(f"- {collapse_ws(label)}")
        md.append("")
    if difficulty.get("Signals"):
        md.append("Difficulty alignment signals:")
        md.extend(f"- {collapse_ws(item)}" for item in difficulty.get("Signals", []))
        md.append("")
    return "\n".join(md).strip() + "\n"


def build_input_markdown(run, evidence_rel_path):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    function_name = normalize_text(summary.get("FunctionName"))
    md = []
    md.append(f"# Input Quality Evidence: {summary.get('RunId', '')}")
    md.append("")
    md.append("## Run Metadata")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Function", function_name),
                ("Recorded at", summary.get("RecordedAt")),
                ("Status", summary.get("Status")),
                ("Subject", request.get("Subject")),
                ("Source file", request.get("FileName")),
                ("Mode", request.get("Mode")),
                ("Model summary", summary.get("ModelSummary")),
                ("Total cost USD", summary.get("Totals", {}).get("TotalCostUsd")),
                ("Latency ms", summary.get("Totals", {}).get("LatencyMs")),
                ("Evidence file", evidence_rel_path),
            ]
        )
    )
    md.append("")
    if function_name == "extracted-content":
        extracted = response.get("Extracted", {}) or {}
        md.append("## Parsed Extracted Result")
        md.append("")
        md.append(
            format_key_value_lines(
                [
                    ("Parser", extracted.get("ParserName")),
                    ("Vision model", extracted.get("VisionModelName")),
                    ("Extraction mode", extracted.get("ExtractionMode")),
                    ("Word count", extracted.get("WordCount")),
                    ("Embedded image count", extracted.get("EmbeddedImageCount")),
                    ("Detected image refs", join_values(extracted.get("DetectedImageReferences", []), sep=", ", limit=6)),
                    ("Warnings", join_values(extracted.get("Warnings", []), sep=" | ")),
                ]
            )
        )
        md.append("")
        md.append("### Normalized Markdown Preview")
        md.append("")
        md.append(shorten(extracted.get("NormalizedMarkdown"), 3500))
        md.append("")
    elif function_name == "embedding-tagging":
        chunks = response.get("Chunks", []) or []
        md.append("## Chunking and Tagging Result")
        md.append("")
        md.append(
            format_key_value_lines(
                [
                    ("Document id", response.get("DocumentId")),
                    ("Chunk count", len(chunks)),
                    ("Embedding model", request.get("EmbeddingModel")),
                    ("Tagging model", request.get("TaggingModel")),
                    ("Allowed tags", join_values(request.get("AllowedTags", []), sep=", ", limit=12)),
                ]
            )
        )
        md.append("")
        for idx, chunk in enumerate(chunks[:5], start=1):
            md.append(f"### Chunk {idx}: {collapse_ws(chunk.get('SectionTitle') or chunk.get('Id'))}")
            md.append("")
            md.append(
                format_key_value_lines(
                    [
                        ("Chunk id", chunk.get("Id")),
                        ("Chunk type", chunk.get("ChunkType")),
                        ("Strategy", chunk.get("ChunkingStrategy")),
                        ("Source pages", f"{normalize_text(chunk.get('SourcePageFrom'))}-{normalize_text(chunk.get('SourcePageTo'))}"),
                        ("Topic tags", join_values(chunk.get("TopicTags", []), sep=", ")),
                        ("Token count", chunk.get("TokenCount")),
                    ]
                )
            )
            md.append("")
            md.append(shorten(chunk.get("MarkdownText") or chunk.get("NormalizedText") or chunk.get("RawText"), 1200))
            md.append("")
    return "\n".join(md).strip() + "\n"


def build_mentor_markdown(run, evidence_rel_path):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    feedback = response.get("Feedback", {}) or {}
    quality = feedback.get("QualityScore", {}) or {}
    performance = feedback.get("PerformanceSummary", {}) or {}
    files = request.get("SourceFiles", []) or []
    md = []
    md.append(f"# Code Mentor Evidence: {summary.get('RunId', '')}")
    md.append("")
    md.append("## Run Metadata")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Recorded at", summary.get("RecordedAt")),
                ("Subject", request.get("Subject")),
                ("Language", request.get("Language")),
                ("Mentor model", request.get("MentorModel")),
                ("Question type", feedback.get("QuestionType")),
                ("Verdict", feedback.get("Verdict")),
                ("Quality score", quality.get("Overall")),
                ("Confidence", quality.get("Confidence")),
                ("Suggested complexity", feedback.get("SuggestedComplexity")),
                ("Total cost USD", summary.get("Totals", {}).get("TotalCostUsd")),
                ("Latency ms", summary.get("Totals", {}).get("LatencyMs")),
                ("Evidence file", evidence_rel_path),
            ]
        )
    )
    md.append("")
    md.append("## Problem Statement")
    md.append("")
    md.append(shorten(request.get("Problem"), 1800))
    md.append("")
    if files:
        md.append("## Source Files")
        md.append("")
        for file in files[:6]:
            md.append(f"- {collapse_ws(file.get('Filename') or file.get('filename'))}: {shorten(file.get('Content') or file.get('content'), 500)}")
        md.append("")
    md.append("## Mentor Feedback")
    md.append("")
    md.append(
        format_key_value_lines(
            [
                ("Issue categories", join_values(feedback.get("IssueCategories", []), sep=", ")),
                ("Feedback text", feedback.get("FeedbackText")),
                ("Performance summary", performance.get("Summary")),
                ("Time complexity", performance.get("TimeComplexity")),
                ("Space complexity", performance.get("SpaceComplexity")),
                ("Performance notes", join_values(performance.get("Notes", []), sep=" | ", limit=12)),
                ("Error analysis", join_values([
                    " | ".join(filter(None, [
                        normalize_text(item.get("Category")),
                        normalize_text(item.get("Severity")),
                        normalize_text(item.get("Title")),
                        normalize_text(item.get("Detail")),
                        join_values(item.get("FailingScenarios", []), sep=", ", limit=6)
                    ]))
                    for item in feedback.get("ErrorAnalysis", [])
                ], sep=" || ", limit=10)),
                ("Improvement suggestions", join_values([
                    " | ".join(filter(None, [
                        normalize_text(item.get("Priority")),
                        normalize_text(item.get("Title")),
                        normalize_text(item.get("Detail")),
                        normalize_text(item.get("ExpectedImpact"))
                    ]))
                    for item in feedback.get("ImprovementSuggestions", [])
                ], sep=" || ", limit=10)),
            ]
        )
    )
    md.append("")
    return "\n".join(md).strip() + "\n"


def build_input_evidence(run):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    function_name = normalize_text(summary.get("FunctionName"))
    run_id = normalize_text(summary.get("RunId"))
    if function_name == "extracted-content":
        extracted = response.get("Extracted", {}) or {}
        preview = shorten(extracted.get("NormalizedMarkdown") or extracted.get("RawText"), 1200)
        sections = []
        for line in normalize_text(extracted.get("NormalizedMarkdown")).splitlines():
            if line.strip().startswith("#"):
                sections.append(line.strip("# ").strip())
        row = {
            "run_id": run_id,
            "function_name": function_name,
            "recorded_at": normalize_text(summary.get("RecordedAt")),
            "subject": normalize_text(request.get("Subject")),
            "source_file": normalize_text(request.get("FileName")),
            "mode": normalize_text(request.get("Mode")),
            "primary_model": normalize_text(request.get("VisionModel") or summary.get("ModelFields", {}).get("VisionModel")),
            "status": normalize_text(summary.get("Status")),
            "input_tokens": normalize_text(summary.get("Totals", {}).get("InputTokens")),
            "output_tokens": normalize_text(summary.get("Totals", {}).get("OutputTokens")),
            "total_cost_usd": normalize_text(summary.get("Totals", {}).get("TotalCostUsd")),
            "latency_ms": normalize_text(summary.get("Totals", {}).get("LatencyMs")),
            "parsed_summary": join_values(
                [
                    f"parser={extracted.get('ParserName')}",
                    f"vision={extracted.get('VisionModelName')}",
                    f"words={extracted.get('WordCount')}",
                    f"images={extracted.get('EmbeddedImageCount')}",
                ],
                sep=" | ",
            ),
            "content_preview": preview,
            "structure_preview": join_values(sections, sep=" | ", limit=8),
            "tags_or_topics": "",
            "human_readable_result": shorten(preview, 500),
            "evidence_file": "",
            "history_file": str(run["path"].relative_to(REPO)).replace("\\", "/"),
        }
        return row, build_input_markdown
    if function_name == "embedding-tagging":
        chunks = response.get("Chunks", []) or []
        tags = []
        titles = []
        previews = []
        strategies = []
        for chunk in chunks:
            tags.extend(chunk.get("TopicTags", []) or [])
            if chunk.get("SectionTitle"):
                titles.append(chunk.get("SectionTitle"))
            if chunk.get("ChunkingStrategy"):
                strategies.append(chunk.get("ChunkingStrategy"))
            if len(previews) < 2:
                previews.append(shorten(chunk.get("MarkdownText") or chunk.get("NormalizedText") or chunk.get("RawText"), 450))
        row = {
            "run_id": run_id,
            "function_name": function_name,
            "recorded_at": normalize_text(summary.get("RecordedAt")),
            "subject": normalize_text(request.get("Subject")),
            "source_file": normalize_text(request.get("FileName")),
            "mode": normalize_text(request.get("Mode")),
            "primary_model": normalize_text(request.get("EmbeddingModel")),
            "status": normalize_text(summary.get("Status")),
            "input_tokens": normalize_text(summary.get("Totals", {}).get("InputTokens")),
            "output_tokens": normalize_text(summary.get("Totals", {}).get("OutputTokens")),
            "total_cost_usd": normalize_text(summary.get("Totals", {}).get("TotalCostUsd")),
            "latency_ms": normalize_text(summary.get("Totals", {}).get("LatencyMs")),
            "parsed_summary": join_values(
                [
                    f"chunks={len(chunks)}",
                    f"embedding={request.get('EmbeddingModel')}",
                    f"tagging={request.get('TaggingModel')}",
                    f"strategies={join_values(sorted(set(strategies)), sep=', ', limit=4)}",
                ],
                sep=" | ",
            ),
            "content_preview": " || ".join(previews),
            "structure_preview": join_values(titles, sep=" | ", limit=8),
            "tags_or_topics": join_values(sorted(set(tags)), sep=", ", limit=15),
            "human_readable_result": f"{len(chunks)} chunks; top sections: {join_values(titles, sep=', ', limit=5)}",
            "evidence_file": "",
            "history_file": str(run["path"].relative_to(REPO)).replace("\\", "/"),
        }
        return row, build_input_markdown
    return None, None


def build_question_evidence(run):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    review = response.get("Review", {}) or {}
    schema = response.get("SchemaMapping", {}) or {}
    difficulty = response.get("DifficultyAlignment", {}) or {}
    chunks = request.get("Chunks", []) or []
    questions = response.get("Questions", []) or []
    row = {
        "run_id": normalize_text(summary.get("RunId")),
        "recorded_at": normalize_text(summary.get("RecordedAt")),
        "subject": normalize_text(request.get("Subject")),
        "question_type": normalize_text(request.get("QuestionType")),
        "difficulty": normalize_text(request.get("Difficulty")),
        "generator_model": normalize_text(request.get("GeneratorModel") or summary.get("ModelFields", {}).get("GeneratorModel")),
        "reviewer_model": normalize_text(request.get("ReviewerModel") or summary.get("ModelFields", {}).get("ReviewerModel")),
        "status": normalize_text(summary.get("Status")),
        "question_count": len(questions),
        "input_chunk_count": len(chunks),
        "input_chunk_refs": join_values([chunk.get("ChunkId") for chunk in chunks], sep=", ", limit=10),
        "source_topic_tags": join_values(sorted({tag for chunk in chunks for tag in (chunk.get("TopicTags") or [])}), sep=", ", limit=16),
        "input_context_preview": join_values([shorten(chunk.get("ContentText"), 220) for chunk in chunks[:3]], sep=" || ", limit=3),
        "generated_titles": join_values([question.get("Title") for question in questions], sep=" | ", limit=6),
        "generated_body_view": "\n\n".join(format_question_block(question, idx) for idx, question in enumerate(questions[:3], start=1)),
        "review_score": normalize_text(review.get("Score")),
        "review_status": normalize_text(review.get("ReviewStatus")),
        "review_issues": join_values(review.get("Issues", []), sep=" | ", limit=10),
        "review_suggestions": join_values(review.get("Suggestions", []), sep=" | ", limit=10),
        "rule_based_summary": join_values([issue.get("Message") for issue in schema.get("Issues", [])] + list(difficulty.get("Signals", []) or []), sep=" | ", limit=12),
        "prompt_versions": "",
        "evidence_file": "",
        "history_file": str(run["path"].relative_to(REPO)).replace("\\", "/"),
    }
    return row, build_question_markdown


def build_mentor_evidence(run):
    summary = run["summary"]
    request = run["request"]
    response = run["response"]
    feedback = response.get("Feedback", {}) or {}
    quality = feedback.get("QualityScore", {}) or {}
    performance = feedback.get("PerformanceSummary", {}) or {}
    files = request.get("SourceFiles", []) or []
    row = {
        "run_id": normalize_text(summary.get("RunId")),
        "recorded_at": normalize_text(summary.get("RecordedAt")),
        "subject": normalize_text(request.get("Subject")),
        "language": normalize_text(request.get("Language")),
        "mentor_model": normalize_text(request.get("MentorModel") or summary.get("ModelFields", {}).get("MentorModel")),
        "status": normalize_text(summary.get("Status")),
        "problem_summary": shorten(request.get("Problem"), 500),
        "source_files": join_values([item.get("Filename") or item.get("filename") for item in files], sep=", ", limit=8),
        "source_code_preview": join_values([shorten(item.get("Content") or item.get("content"), 180) for item in files[:3]], sep=" || ", limit=3),
        "question_type": normalize_text(feedback.get("QuestionType")),
        "mentor_verdict": normalize_text(feedback.get("Verdict")),
        "issue_categories": join_values(feedback.get("IssueCategories", []), sep=", ", limit=8),
        "quality_score_overall": normalize_number(quality.get("Overall"), 4),
        "quality_score_correctness": normalize_number(quality.get("Correctness"), 4),
        "quality_score_robustness": normalize_number(quality.get("Robustness"), 4),
        "quality_score_code_quality": normalize_number(quality.get("CodeQuality"), 4),
        "quality_score_efficiency": normalize_number(quality.get("Efficiency"), 4),
        "quality_confidence": normalize_number(quality.get("Confidence"), 4),
        "feedback_text": normalize_text(feedback.get("FeedbackText")),
        "performance_summary_view": join_values(
            [
                normalize_text(performance.get("Summary")),
                normalize_text(performance.get("TimeComplexity")),
                normalize_text(performance.get("SpaceComplexity")),
                join_values(performance.get("Notes", []), sep=" | ", limit=8),
            ],
            sep=" || ",
            limit=4),
        "error_analysis_view": join_values([
            " | ".join(filter(None, [
                normalize_text(item.get("Category")),
                normalize_text(item.get("Severity")),
                normalize_text(item.get("Title")),
                normalize_text(item.get("Detail"))
            ]))
            for item in feedback.get("ErrorAnalysis", [])
        ], sep=" || ", limit=10),
        "improvement_suggestions_view": join_values([
            " | ".join(filter(None, [
                normalize_text(item.get("Priority")),
                normalize_text(item.get("Title")),
                normalize_text(item.get("Detail"))
            ]))
            for item in feedback.get("ImprovementSuggestions", [])
        ], sep=" || ", limit=10),
        "failing_scenarios_view": join_values([
            join_values(item.get("FailingScenarios", []), sep=", ", limit=6)
            for item in feedback.get("ErrorAnalysis", [])
        ], sep=" || ", limit=10),
        "suggested_complexity_view": normalize_text(feedback.get("SuggestedComplexity") or performance.get("TimeComplexity")),
        "evidence_file": "",
        "history_file": str(run["path"].relative_to(REPO)).replace("\\", "/"),
    }
    return row, build_mentor_markdown


def build_input_benchmark_rows(run_payloads, config_refs_by_run):
    rows = []
    for run in run_payloads:
        summary = run["summary"]
        request = run["request"]
        response = run["response"]
        function_name = normalize_text(summary.get("FunctionName"))
        if function_name not in {"extracted-content", "embedding-tagging"}:
            continue
        refs = config_refs_by_run.get(normalize_text(summary.get("RunId")), [])
        ref_map = get_config_ref_map(refs)
        totals = summary.get("Totals", {}) or {}
        draft = response.get("Draft", {}) or {}
        chunks = response.get("Chunks", []) or []
        extracted = response.get("Extracted", {}) or {}
        input_tokens = normalize_int(totals.get("InputTokens"))
        output_tokens = normalize_int(totals.get("OutputTokens"))
        total_tokens = (input_tokens or 0) + (output_tokens or 0)
        cost_usd = normalize_number(totals.get("TotalCostUsd"), 6)
        latency_ms = normalize_number(totals.get("LatencyMs"), 2)
        avg_chunk_tokens = None
        avg_chunk_words = None
        if chunks:
            token_values = [normalize_number(chunk.get("TokenCount")) for chunk in chunks if normalize_number(chunk.get("TokenCount")) is not None]
            word_values = [normalize_number(chunk.get("WordCount")) for chunk in chunks if normalize_number(chunk.get("WordCount")) is not None]
            avg_chunk_tokens = round(sum(token_values) / len(token_values), 2) if token_values else None
            avg_chunk_words = round(sum(word_values) / len(word_values), 2) if word_values else None

        rows.append(
            {
                "run_id": normalize_text(summary.get("RunId")),
                "test_date": normalize_text(summary.get("RecordedAt"))[:10],
                "tester": DEFAULT_TESTER,
                "function_group": function_name,
                "document_id": normalize_text(request.get("DocumentId") or draft.get("DocumentId") or response.get("DocumentId")),
                "document_name": normalize_text(request.get("FileName") or draft.get("SourceName")),
                "document_type": normalize_text(draft.get("SourceType") or Path(normalize_text(request.get("FileName"))).suffix.lstrip(".")),
                "subject_code": normalize_text(request.get("Subject") or draft.get("SubjectCode")),
                "syllabus_scope": normalize_text(draft.get("OwnershipType")),
                "source_pages": normalize_text(draft.get("TotalPages")),
                "source_word_count": normalize_int(extracted.get("WordCount") or draft.get("PricingBasisSnapshot", {}).get("TotalWordsEstimate")),
                "source_image_count": normalize_int(extracted.get("EmbeddedImageCount") or draft.get("EmbeddedImageCount")),
                "pipeline_mode": normalize_text(request.get("Mode") or extracted.get("ExtractionMode") or draft.get("ExtractionMode")),
                "ocr_model": normalize_text(request.get("VisionModel") or extracted.get("VisionModelName") or draft.get("VisionModel")),
                "embedding_model": normalize_text(request.get("EmbeddingModel")),
                "tagging_model": normalize_text(request.get("TaggingModel")),
                "chunking_strategy": normalize_text(request.get("ChunkingStrategy") or (chunks[0].get("ChunkingStrategy") if chunks else "")),
                "human_review_used": yes_no(draft.get("ReviewedAt") or draft.get("ApprovalVersion")),
                "prompt_version": get_config_version(refs, "prompt"),
                "policy_version": get_config_version(refs, "policy"),
                "rubric_version": get_config_version(refs, "rubric"),
                "ground_truth_set_id": get_config_version(refs, "groundtruth"),
                "input_tokens": input_tokens,
                "output_tokens": output_tokens,
                "vision_input_tokens": input_tokens if function_name == "extracted-content" else None,
                "vision_output_tokens": output_tokens if function_name == "extracted-content" else None,
                "total_tokens": total_tokens if total_tokens else None,
                "latency_ms": latency_ms,
                "cost_usd": cost_usd,
                "num_chunks": normalize_int(len(chunks)) if chunks else None,
                "avg_chunk_tokens": avg_chunk_tokens,
                "avg_chunk_words": avg_chunk_words,
                "tag_coverage_score": "",
                "tag_relevance_score": "",
                "chunk_boundary_score": "",
                "content_cleanliness_score": "",
                "ocr_text_accuracy_score": "",
                "image_description_quality_score": "",
                "grounding_traceability_score": "",
                "overall_quality_score": "",
                "quality_cost_index": "",
                "quality_latency_index": "",
                "pass_fail": "PASS" if normalize_text(summary.get("Status")) not in {"failed", "error"} else "FAIL",
                "main_issues": join_values(
                    extracted.get("Warnings", [])
                    or draft.get("CleanupWarnings", [])
                    or ([normalize_text(draft.get("ReviewStatus"))] if normalize_text(draft.get("ReviewStatus")) not in {"", "approved"} else []),
                    sep=" | ",
                    limit=12,
                ),
                "notes": join_values(
                    [
                        normalize_text(ref_map.get("policy", {}).get("change_summary")),
                        f"mode={normalize_text(request.get('Mode') or extracted.get('ExtractionMode') or draft.get('ExtractionMode'))}" if normalize_text(request.get("Mode") or extracted.get("ExtractionMode") or draft.get("ExtractionMode")) else "",
                        f"parser={normalize_text(extracted.get('ParserName') or draft.get('IngestParser'))}" if normalize_text(extracted.get("ParserName") or draft.get("IngestParser")) else "",
                        f"vision={normalize_text(request.get('VisionModel') or extracted.get('VisionModelName') or draft.get('VisionModel'))}" if normalize_text(request.get("VisionModel") or extracted.get("VisionModelName") or draft.get("VisionModel")) else "",
                        f"embedding={normalize_text(request.get('EmbeddingModel'))}" if normalize_text(request.get("EmbeddingModel")) else "",
                        f"tagging={normalize_text(request.get('TaggingModel'))}" if normalize_text(request.get("TaggingModel")) else "",
                        f"chunking={normalize_text(request.get('ChunkingStrategy') or (chunks[0].get('ChunkingStrategy') if chunks else ''))}" if normalize_text(request.get("ChunkingStrategy") or (chunks[0].get("ChunkingStrategy") if chunks else "")) else "",
                    ],
                    sep=" | ",
                    limit=8,
                ),
            }
        )
    rows.sort(key=lambda item: parse_dt(item["test_date"]))
    return rows


def build_question_benchmark_rows(run_payloads, config_refs_by_run, active_prompt_versions):
    rows = []
    for run in run_payloads:
        summary = run["summary"]
        request = run["request"]
        response = run["response"]
        if normalize_text(summary.get("FunctionName")) != "generation-review":
            continue
        refs = config_refs_by_run.get(normalize_text(summary.get("RunId")), [])
        totals = summary.get("Totals", {}) or {}
        review = response.get("Review", {}) or {}
        schema = response.get("SchemaMapping", {}) or {}
        difficulty = response.get("DifficultyAlignment", {}) or {}
        questions = response.get("Questions", []) or []
        chunks = request.get("Chunks", []) or []
        generator_costs = get_stage_costs(response, "question-generation")
        reviewer_costs = get_stage_costs(response, "question-review")
        generator_input_tokens = sum(normalize_int(cost.get("InputTokens")) or 0 for cost in generator_costs) or None
        generator_output_tokens = sum(normalize_int(cost.get("OutputTokens")) or 0 for cost in generator_costs) or None
        reviewer_input_tokens = sum(normalize_int(cost.get("InputTokens")) or 0 for cost in reviewer_costs) or None
        reviewer_output_tokens = sum(normalize_int(cost.get("OutputTokens")) or 0 for cost in reviewer_costs) or None
        generator_cost_usd = round(sum(normalize_number(cost.get("TotalCostUsd")) or 0 for cost in generator_costs), 6) if generator_costs else None
        reviewer_cost_usd = round(sum(normalize_number(cost.get("TotalCostUsd")) or 0 for cost in reviewer_costs), 6) if reviewer_costs else None
        total_tokens = (normalize_int(totals.get("InputTokens")) or 0) + (normalize_int(totals.get("OutputTokens")) or 0)
        review_score = normalize_number(review.get("Score"), 4)
        input_chunk_token_count = (
            sum(
                (normalize_int(chunk.get("TokenCount")) if normalize_int(chunk.get("TokenCount")) is not None else estimate_token_count(chunk.get("ContentText")))
                for chunk in chunks
            ) or None
        )
        accepted_count = len(questions) if normalize_text(review.get("NeedsRevision")).lower() not in {"true", "1"} and normalize_text(summary.get("Status")) == "accepted" else 0
        rejected_count = max(len(questions) - accepted_count, 0)
        schema_valid_rate = 1.0 if schema.get("IsValid") else 0.0 if schema else None
        grounded_rate = 1.0 if review.get("ContentGrounded") else 0.0 if review else None
        difficulty_alignment_score = 1.0 if difficulty.get("IsAligned") else review_score if review_score is not None else 0.0
        rule_issues = schema.get("Issues", []) or []
        rule_based_pass_rate = 1.0 if not rule_issues else 0.0
        final_quality_score = review_score
        total_cost_usd = normalize_number(totals.get("TotalCostUsd"), 6)
        latency_ms = normalize_number(totals.get("LatencyMs"), 2)
        context_pack_id = normalize_text(request.get("ContextPackId") or response.get("ContextPackId"))
        context_source = "context_pack" if context_pack_id else "chunks"
        retrieval_strategy = normalize_text(request.get("RetrievalStrategy")) or ("context_pack" if context_pack_id else "direct_chunks")
        compression_base = generator_input_tokens or normalize_int(totals.get("InputTokens"))
        prompt_version_generator = get_config_version([ref for ref in refs if ref.get("artifact_type") == "prompt" and "question_generation" in ref.get("stable_id", "")], "prompt", default="")
        prompt_version_reviewer = get_config_version([ref for ref in refs if ref.get("artifact_type") == "prompt" and "question_review" in ref.get("stable_id", "")], "prompt", default="")
        if not prompt_version_generator:
            current = active_prompt_versions.get("question_generation")
            prompt_version_generator = f"question_generation:{current}" if current else "N/A"
        if not prompt_version_reviewer:
            current = active_prompt_versions.get("question_review")
            prompt_version_reviewer = f"question_review:{current}" if current else "N/A"

        rows.append(
            {
                "run_id": normalize_text(summary.get("RunId")),
                "test_date": normalize_text(summary.get("RecordedAt"))[:10],
                "tester": DEFAULT_TESTER,
                "subject_code": normalize_text(request.get("Subject")),
                "question_type": normalize_text(request.get("QuestionType")),
                "difficulty_target": normalize_text(request.get("Difficulty")),
                "generation_mode": normalize_text(request.get("Mode")),
                "generator_model": normalize_text(request.get("GeneratorModel") or summary.get("ModelFields", {}).get("GeneratorModel")),
                "reviewer_model": normalize_text(request.get("ReviewerModel") or summary.get("ModelFields", {}).get("ReviewerModel")),
                "dual_agent_mode": "yes" if normalize_text(request.get("Mode")) == "DualAgent" else "no",
                "context_source": context_source,
                "context_pack_id": context_pack_id,
                "knowledge_chunk_count": len(chunks),
                "packed_token_count": input_chunk_token_count,
                "compression_ratio": safe_ratio(input_chunk_token_count, compression_base),
                "retrieval_strategy": retrieval_strategy,
                "prompt_version_generator": prompt_version_generator,
                "prompt_version_reviewer": prompt_version_reviewer,
                "rubric_version": get_config_version(refs, "rubric"),
                "policy_version": get_config_version(refs, "policy"),
                "ground_truth_set_id": get_config_version(refs, "groundtruth"),
                "requested_question_count": normalize_int(request.get("Count")),
                "generated_question_count": len(questions),
                "accepted_question_count": accepted_count,
                "rejected_question_count": rejected_count,
                "revision_rounds": normalize_int(response.get("AttemptsUsed")),
                "generator_input_tokens": generator_input_tokens,
                "generator_output_tokens": generator_output_tokens,
                "reviewer_input_tokens": reviewer_input_tokens,
                "reviewer_output_tokens": reviewer_output_tokens,
                "total_tokens": total_tokens if total_tokens else None,
                "generator_cost_usd": generator_cost_usd,
                "reviewer_cost_usd": reviewer_cost_usd,
                "total_cost_usd": total_cost_usd,
                "latency_ms": latency_ms,
                "schema_valid_rate": schema_valid_rate,
                "grounded_rate": grounded_rate,
                "difficulty_alignment_score": difficulty_alignment_score,
                "topic_relevance_score": "",
                "question_clarity_score": "",
                "question_usefulness_score": "",
                "distractor_quality_score": "" if normalize_text(request.get("QuestionType")) != "FE" else "",
                "testcase_quality_score": "" if normalize_text(request.get("QuestionType")) != "PE" else "",
                "rule_based_pass_rate": rule_based_pass_rate,
                "final_quality_score": final_quality_score,
                "quality_cost_index": score_cost_index(final_quality_score, total_cost_usd),
                "quality_latency_index": score_latency_index(final_quality_score, latency_ms),
                "pass_fail": "PASS" if normalize_text(summary.get("Status")) in {"accepted", "completed"} else "FAIL",
                "failure_reason": join_values(review.get("Issues", []), sep=" | ", limit=10),
                "notes": join_values(review.get("Suggestions", []), sep=" | ", limit=10),
            }
        )
    rows.sort(key=lambda item: parse_dt(item["test_date"]))
    return rows


def build_mentor_benchmark_rows(run_payloads, config_refs_by_run):
    rows = []
    for run in run_payloads:
        summary = run["summary"]
        request = run["request"]
        response = run["response"]
        if normalize_text(summary.get("FunctionName")) != "code-mentor":
            continue
        refs = config_refs_by_run.get(normalize_text(summary.get("RunId")), [])
        totals = summary.get("Totals", {}) or {}
        feedback = response.get("Feedback", {}) or {}
        quality = feedback.get("QualityScore", {}) or {}
        files = request.get("SourceFiles", []) or []
        input_tokens = normalize_int(totals.get("InputTokens"))
        output_tokens = normalize_int(totals.get("OutputTokens"))
        total_tokens = (input_tokens or 0) + (output_tokens or 0)
        overall_quality = normalize_number(quality.get("Overall"), 4)
        cost_usd = normalize_number(totals.get("TotalCostUsd"), 6)
        latency_ms = normalize_number(totals.get("LatencyMs"), 2)
        rows.append(
            {
                "run_id": normalize_text(summary.get("RunId")),
                "test_date": normalize_text(summary.get("RecordedAt"))[:10],
                "tester": DEFAULT_TESTER,
                "subject_code": normalize_text(request.get("Subject")),
                "language": normalize_text(request.get("Language")),
                "exercise_type": normalize_text(feedback.get("QuestionType") or "PE"),
                "submission_id": normalize_text(request.get("SubmissionId")),
                "submission_size": len(normalize_text(request.get("Code"))) if request.get("Code") else sum(len(normalize_text(item.get("Content") or item.get("content"))) for item in files),
                "file_count": len(files) if files else (1 if request.get("Code") else 0),
                "model": normalize_text(request.get("MentorModel") or summary.get("ModelFields", {}).get("MentorModel")),
                "prompt_version": get_config_version(refs, "prompt"),
                "rubric_version": get_config_version(refs, "rubric"),
                "policy_version": get_config_version(refs, "policy"),
                "ground_truth_set_id": get_config_version(refs, "groundtruth"),
                "input_tokens": input_tokens,
                "output_tokens": output_tokens,
                "total_tokens": total_tokens if total_tokens else None,
                "latency_ms": latency_ms,
                "cost_usd": cost_usd,
                "bug_detection_score": normalize_number(quality.get("Correctness"), 4),
                "correctness_feedback_score": normalize_number(quality.get("Correctness"), 4),
                "code_quality_feedback_score": normalize_number(quality.get("CodeQuality"), 4),
                "actionability_score": "",
                "specificity_score": "",
                "hallucination_penalty": "",
                "grounded_to_code_score": normalize_number(quality.get("Robustness"), 4),
                "pedagogical_value_score": normalize_number(quality.get("Overall"), 4),
                "overall_quality_score": overall_quality,
                "quality_cost_index": score_cost_index(overall_quality, cost_usd),
                "quality_latency_index": score_latency_index(overall_quality, latency_ms),
                "expected_feedback_match": "",
                "pass_fail": "PASS" if normalize_text(feedback.get("Verdict")) not in {"failed", "error"} else "FAIL",
                "main_issues": join_values([item.get("Title") for item in feedback.get("ErrorAnalysis", [])], sep=" | ", limit=10),
                "notes": normalize_text(feedback.get("FeedbackText")),
            }
        )
    rows.sort(key=lambda item: parse_dt(item["test_date"]))
    return rows


def generate_evidence(run_payloads, config_refs_by_run):
    input_rows = []
    question_rows = []
    mentor_rows = []
    for run in run_payloads:
        function_name = normalize_text(run["summary"].get("FunctionName"))
        run_id = normalize_text(run["summary"].get("RunId"))
        prompt_refs = [f"{ref['stable_id']}:{ref['version']}" for ref in config_refs_by_run.get(run_id, []) if ref["artifact_type"] == "prompt"]
        if function_name in {"extracted-content", "embedding-tagging"}:
            row, markdown_builder = build_input_evidence(run)
            if row:
                evidence_path = EVIDENCE_ROOT / "input-quality" / f"{run_id}.md"
                evidence_rel = str(evidence_path.relative_to(REPO)).replace("\\", "/")
                row["evidence_file"] = evidence_rel
                write_text(evidence_path, markdown_builder(run, evidence_rel))
                input_rows.append(row)
        elif function_name == "generation-review":
            row, markdown_builder = build_question_evidence(run)
            evidence_path = EVIDENCE_ROOT / "question-generation" / f"{run_id}.md"
            evidence_rel = str(evidence_path.relative_to(REPO)).replace("\\", "/")
            row["prompt_versions"] = " | ".join(prompt_refs)
            row["evidence_file"] = evidence_rel
            write_text(evidence_path, markdown_builder(run, evidence_rel))
            question_rows.append(row)
        elif function_name == "code-mentor":
            row, markdown_builder = build_mentor_evidence(run)
            evidence_path = EVIDENCE_ROOT / "code-mentor" / f"{run_id}.md"
            evidence_rel = str(evidence_path.relative_to(REPO)).replace("\\", "/")
            row["evidence_file"] = evidence_rel
            write_text(evidence_path, markdown_builder(run, evidence_rel))
            mentor_rows.append(row)
    input_rows.sort(key=lambda item: parse_dt(item["recorded_at"]))
    question_rows.sort(key=lambda item: parse_dt(item["recorded_at"]))
    mentor_rows.sort(key=lambda item: parse_dt(item["recorded_at"]))
    return input_rows, question_rows, mentor_rows


def main():
    ensure_required_paths()
    bootstrap_output_structure()

    if not WORKBOOK_PATH.exists():
        shutil.copy2(WORKBOOK_TEMPLATE_PATH, WORKBOOK_PATH)
    changelog_seed = REPORT_DIR / "ai_config_version_changelog.csv"
    if not CHANGELOG_CSV.exists() and changelog_seed.exists():
        shutil.copy2(changelog_seed, CHANGELOG_CSV)

    snapshot_rows, snapshot_index = load_snapshot_index()
    catalog_rows = load_current_catalog()
    active_prompt_versions = load_active_prompt_versions()
    changelog_rows = load_changelog_rows()
    version_tracking_rows = build_version_tracking_rows(changelog_rows)
    prompt_compare_rows = build_prompt_comparison_rows(snapshot_rows)
    run_index_rows, config_ref_rows, config_refs_by_run, run_payloads = collect_runs(snapshot_index)
    input_evidence_rows, question_evidence_rows, mentor_evidence_rows = generate_evidence(run_payloads, config_refs_by_run)
    input_benchmark_rows = build_input_benchmark_rows(run_payloads, config_refs_by_run)
    question_benchmark_rows = build_question_benchmark_rows(run_payloads, config_refs_by_run, active_prompt_versions)
    mentor_benchmark_rows = build_mentor_benchmark_rows(run_payloads, config_refs_by_run)

    input_benchmark_headers = ["run_id", "test_date", "tester", "function_group", "document_id", "document_name", "document_type", "subject_code", "syllabus_scope", "source_pages", "source_word_count", "source_image_count", "pipeline_mode", "ocr_model", "embedding_model", "tagging_model", "chunking_strategy", "human_review_used", "prompt_version", "policy_version", "rubric_version", "ground_truth_set_id", "input_tokens", "output_tokens", "vision_input_tokens", "vision_output_tokens", "total_tokens", "latency_ms", "cost_usd", "num_chunks", "avg_chunk_tokens", "avg_chunk_words", "tag_coverage_score", "tag_relevance_score", "chunk_boundary_score", "content_cleanliness_score", "ocr_text_accuracy_score", "image_description_quality_score", "grounding_traceability_score", "overall_quality_score", "quality_cost_index", "quality_latency_index", "pass_fail", "main_issues", "notes"]
    question_benchmark_headers = ["run_id", "test_date", "tester", "subject_code", "question_type", "difficulty_target", "generation_mode", "generator_model", "reviewer_model", "dual_agent_mode", "context_source", "context_pack_id", "knowledge_chunk_count", "packed_token_count", "compression_ratio", "retrieval_strategy", "prompt_version_generator", "prompt_version_reviewer", "rubric_version", "policy_version", "ground_truth_set_id", "requested_question_count", "generated_question_count", "accepted_question_count", "rejected_question_count", "revision_rounds", "generator_input_tokens", "generator_output_tokens", "reviewer_input_tokens", "reviewer_output_tokens", "total_tokens", "generator_cost_usd", "reviewer_cost_usd", "total_cost_usd", "latency_ms", "schema_valid_rate", "grounded_rate", "difficulty_alignment_score", "topic_relevance_score", "question_clarity_score", "question_usefulness_score", "distractor_quality_score", "testcase_quality_score", "rule_based_pass_rate", "final_quality_score", "quality_cost_index", "quality_latency_index", "pass_fail", "failure_reason", "notes"]
    mentor_benchmark_headers = ["run_id", "test_date", "tester", "subject_code", "language", "exercise_type", "submission_id", "submission_size", "file_count", "model", "prompt_version", "rubric_version", "policy_version", "ground_truth_set_id", "input_tokens", "output_tokens", "total_tokens", "latency_ms", "cost_usd", "bug_detection_score", "correctness_feedback_score", "code_quality_feedback_score", "actionability_score", "specificity_score", "hallucination_penalty", "grounded_to_code_score", "pedagogical_value_score", "overall_quality_score", "quality_cost_index", "quality_latency_index", "expected_feedback_match", "pass_fail", "main_issues", "notes"]
    run_config_headers = ["run_id", "function_name", "route_key", "recorded_at", "artifact_type", "stable_id", "version", "subject_code", "question_type", "language", "content_hash", "source", "source_commit", "change_reason", "change_summary", "snapshot_file", "history_file"]
    run_index_headers = ["run_id", "function_name", "route_key", "recorded_at", "status", "model_summary", "subject", "question_type", "difficulty", "input_tokens", "output_tokens", "total_cost_usd", "latency_ms", "prompt_key_guess", "prompt_version_ref", "config_link_mode", "history_file", "usage_source", "error"]
    input_evidence_headers = ["run_id", "function_name", "recorded_at", "subject", "source_file", "mode", "primary_model", "status", "input_tokens", "output_tokens", "total_cost_usd", "latency_ms", "parsed_summary", "content_preview", "structure_preview", "tags_or_topics", "human_readable_result", "evidence_file", "history_file"]
    question_evidence_headers = ["run_id", "recorded_at", "subject", "question_type", "difficulty", "generator_model", "reviewer_model", "status", "question_count", "input_chunk_count", "input_chunk_refs", "source_topic_tags", "input_context_preview", "generated_titles", "generated_body_view", "review_score", "review_status", "review_issues", "review_suggestions", "rule_based_summary", "prompt_versions", "evidence_file", "history_file"]
    mentor_evidence_headers = ["run_id", "recorded_at", "subject", "language", "mentor_model", "status", "problem_summary", "source_files", "source_code_preview", "question_type", "mentor_verdict", "issue_categories", "quality_score_overall", "quality_score_correctness", "quality_score_robustness", "quality_score_code_quality", "quality_score_efficiency", "quality_confidence", "feedback_text", "performance_summary_view", "error_analysis_view", "improvement_suggestions_view", "failing_scenarios_view", "suggested_complexity_view", "evidence_file", "history_file"]

    write_csv(INPUT_BENCHMARK_CSV, input_benchmark_headers, input_benchmark_rows)
    write_csv(QUESTION_BENCHMARK_CSV, question_benchmark_headers, question_benchmark_rows)
    write_csv(MENTOR_BENCHMARK_CSV, mentor_benchmark_headers, mentor_benchmark_rows)
    write_csv(RUN_CONFIG_CSV, run_config_headers, config_ref_rows)
    write_csv(RUN_INDEX_CSV, run_index_headers, run_index_rows)
    write_csv(INPUT_EVIDENCE_CSV, input_evidence_headers, input_evidence_rows)
    write_csv(QUESTION_EVIDENCE_CSV, question_evidence_headers, question_evidence_rows)
    write_csv(MENTOR_EVIDENCE_CSV, mentor_evidence_headers, mentor_evidence_rows)

    workbook = load_workbook(WORKBOOK_PATH)
    sheet_specs = {
        "Input Quality Benchmark": (input_benchmark_headers, input_benchmark_rows),
        "Question Generation Benchmark": (question_benchmark_headers, question_benchmark_rows),
        "Code Mentor Benchmark": (mentor_benchmark_headers, mentor_benchmark_rows),
        "AI Config Version Tracking": (["config_type", "function_name", "component_name", "version_id", "parent_version_id", "status", "applied_to_models", "subject_code", "question_type", "effective_from", "effective_to", "change_reason", "change_summary", "expected_impact", "dataset_used", "benchmark_run_range", "quality_before", "quality_after", "cost_before", "cost_after", "stability_before", "stability_after", "owner_note", "source_commit", "snapshot_path"], version_tracking_rows),
        "AI Config Catalog": (["config_type", "config_key", "version", "name_or_description", "subject_code", "question_type", "status", "system_prompt_len", "user_prompt_len", "updated_at", "source_file"], catalog_rows),
        "AI Config Snapshots": (["artifact_type", "stable_id", "version", "function_name", "subject_code", "question_type", "captured_at", "source", "source_commit", "captured_from", "change_reason", "change_summary", "is_active", "content_hash", "snapshot_file"], snapshot_rows),
        "Prompt Version Comparison": (["prompt_key", "current_version", "previous_version", "description", "current_system_len", "previous_system_len", "current_user_len", "previous_user_len", "change_detected", "change_reason", "dataset_used", "benchmark_run_ids_before", "benchmark_run_ids_after", "quality_before", "quality_after", "cost_before", "cost_after", "decision_note", "source_commit_current", "source_commit_previous", "current_hash", "previous_hash"], prompt_compare_rows),
        "Run History Index": (run_index_headers, run_index_rows),
        "Run Config References": (run_config_headers, config_ref_rows),
        "Input Quality Evidence": (input_evidence_headers, input_evidence_rows),
        "Question Generation Evidence": (question_evidence_headers, question_evidence_rows),
        "Code Mentor Evidence": (mentor_evidence_headers, mentor_evidence_rows),
    }
    for sheet_name, (headers, rows) in sheet_specs.items():
        worksheet = ensure_sheet(workbook, sheet_name)
        write_rows(worksheet, headers, rows)
        set_tab_style(worksheet)
    workbook.save(WORKBOOK_PATH)

    print(f"Updated workbook: {WORKBOOK_PATH}")
    print(f"Wrote CSV: {INPUT_BENCHMARK_CSV}")
    print(f"Wrote CSV: {QUESTION_BENCHMARK_CSV}")
    print(f"Wrote CSV: {MENTOR_BENCHMARK_CSV}")
    print(f"Wrote CSV: {RUN_CONFIG_CSV}")
    print(f"Wrote CSV: {RUN_INDEX_CSV}")
    print(f"Wrote CSV: {INPUT_EVIDENCE_CSV}")
    print(f"Wrote CSV: {QUESTION_EVIDENCE_CSV}")
    print(f"Wrote CSV: {MENTOR_EVIDENCE_CSV}")
    print(f"Wrote evidence markdown under: {EVIDENCE_ROOT}")


if __name__ == "__main__":
    main()
